"""Noninteractive checks executed by the actual frozen application."""
import json
import os
import tempfile
import traceback
import socket
import subprocess
import time
from pathlib import Path

from .paths import ROOT


def run():
    report = {"root": str(ROOT), "ok": False}
    server = None
    try:
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        import pyaudio
        import win32clipboard
        import win32com.client
        from pycaw.pycaw import AudioUtilities
        from vosk import Model, SetLogLevel
        from .speech import vosk_path, available_voices
        from .qt import QApplication
        from .storage import Store
        from .ui import Window
        from .office import Office
        import requests

        app = QApplication([])
        from .apps import AppCatalog
        from .system_actions import brightness, configured_modes, screenshot
        catalog = AppCatalog()
        catalog.refresh()
        catalog.thread.join(30)
        report["applications"] = len(catalog.entries)
        assert catalog.entries, catalog.error
        report["brightness"] = brightness()
        report["power_modes"] = configured_modes()
        from .hotkey import GlobalHotkey
        hotkey = GlobalHotkey()
        hotkey.start()
        try:
            assert hotkey.ready.wait(5), "Global keyboard hook did not start"
            report["global_hotkey"] = True
        finally:
            hotkey.stop()
            hotkey.thread.join(5)
        assert not hotkey.thread.is_alive(), "Keyboard hook did not stop"
        with tempfile.TemporaryDirectory() as temp:
            store = Store(Path(temp))
            store.config["speak"] = False
            store.config["search_roots"] = [temp]
            shot = Path(temp) / "screenshot.png"
            screenshot(shot)
            assert shot.stat().st_size > 100
            report["screenshot"] = True
            window = Window(store)
            window.show()
            app.processEvents()
            window.close()
            window.speaker.thread.join(3)
            window.engine.file_index.thread.join(10)
            from docx import Document
            from openpyxl import Workbook, load_workbook
            doc = Document()
            doc.add_paragraph("Проверка")
            doc.save(Path(temp) / "test.docx")
            book = Workbook()
            book.active["A1"] = 42
            book.save(Path(temp) / "test.xlsx")
            assert load_workbook(Path(temp) / "test.xlsx").active["A1"].value == 42
        SetLogLevel(-1)
        model = Model(vosk_path(ROOT / "models/vosk-model-small-ru-0.22"))
        report["vosk"] = bool(model)
        report["voices"] = len(available_voices())
        audio = pyaudio.PyAudio()
        report["audio_devices"] = audio.get_device_count()
        audio.terminate()
        assert (ROOT / "runtime/ollama/ollama.exe").is_file()
        config = Store().config
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
        env = os.environ.copy()
        env["OLLAMA_HOST"] = f"127.0.0.1:{port}"
        env["OLLAMA_MODELS"] = str(ROOT / "models/ollama")
        server = subprocess.Popen([str(ROOT / "runtime/ollama/ollama.exe"), "serve"], env=env,
            stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            creationflags=subprocess.CREATE_NO_WINDOW)
        url = f"http://127.0.0.1:{port}"
        for _ in range(80):
            try:
                requests.get(url + "/api/tags", timeout=1).raise_for_status()
                break
            except requests.RequestException:
                time.sleep(.25)
        response = requests.post(url + "/api/chat", json={
            "model": config["ollama_model"], "stream": False,
            "keep_alive": 0,
            "messages": [{"role": "user", "content": "Ответь одним словом: столица Франции?"}],
            "options": {"num_predict": 24, "num_ctx": 2048}}, timeout=180)
        response.raise_for_status()
        report["ai_answer"] = response.json()["message"]["content"]
        assert report["ai_answer"].strip()
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    finally:
        if server is not None:
            server.terminate()
            server.wait(timeout=10)
    (ROOT / "self-test.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1
