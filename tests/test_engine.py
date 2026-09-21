import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from friday.engine import Engine, duration
from friday.office import cell_address
from friday.storage import Store


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = Store(Path(self.temp.name))
        self.windows = Mock()
        self.now = 1800000000.0
        self.engine = Engine(self.store, self.windows, lambda: self.now)

    def test_power_requires_explicit_confirmation_and_expires(self):
        self.assertIn("подтверждаю", self.engine.execute("Пятница, выключи компьютер"))
        self.windows.power.assert_not_called()
        self.engine.execute("да")
        self.windows.power.assert_not_called()
        self.engine.execute("отмена")
        self.engine.execute("подтверждаю")
        self.windows.power.assert_not_called()
        self.engine.execute("выключи компьютер")
        self.now += 31
        self.assertIn("истекло", self.engine.execute("подтверждаю"))
        self.windows.power.assert_not_called()
        self.engine.execute("выключи компьютер")
        self.engine.execute("подтверждаю")
        self.windows.power.assert_called_once_with("shutdown")

    def test_app_alias_case_and_wake_prefix(self):
        self.engine.execute("ПЯТНИЦА, Открой Word!")
        self.windows.open_app.assert_called_once_with("word", self.store.config)

    def test_volume_bounds_and_mute_are_not_toggle(self):
        self.assertIn("от 0 до 100", self.engine.execute("громкость 101"))
        self.windows.volume.assert_not_called()
        self.engine.execute("громкость 30 процентов")
        self.windows.volume.assert_called_once_with(30)
        self.engine.execute("включи звук")
        self.windows.media.assert_called_once_with("unmute")

    def test_reminders_survive_restart_and_fire_only_once(self):
        self.engine.execute("через 5 минут напомни сделать перерыв")
        other = Engine(Store(Path(self.temp.name)), self.windows, lambda: self.now)
        self.assertEqual(len(other.reminders()), 1)
        self.assertEqual(other.due_reminders(), [])
        self.now += 301
        self.assertEqual(other.due_reminders()[0]["text"], "сделать перерыв")
        self.assertEqual(other.due_reminders(), [])

    def test_alternative_reminder_word_order(self):
        result = self.engine.execute("напомни через пять минут позвонить")
        self.assertIn("Напомню", result)
        self.assertEqual(self.engine.reminders()[0]["due"], self.now + 300)

    def test_compound_duration_and_bad_input(self):
        self.assertEqual(duration("1 час 20 минут"), 4800)
        self.assertEqual(duration("пять минут"), 300)
        self.assertEqual(duration("двадцать одна минута"), 1260)
        for value in ("-5 минут", "0 минут", "5", "999999 дней", "не пять минут"):
            with self.assertRaises(ValueError):
                duration(value)

    def test_spoken_numbers_and_cells(self):
        self.engine.execute("громкость тридцать два процента")
        self.engine.execute("громкость тридцать два процентов")
        self.windows.volume.assert_called_with(32)
        self.engine.execute("создай таблицу")
        self.assertIn("Записано", self.engine.execute("запиши одна тысяча пятьсот двадцать три в бэ три"))
        self.assertEqual(self.engine.office.read("B3"), "1523")

    def test_dictation_persists_and_saves_without_overwrite(self):
        self.engine.execute("начни диктовку")
        self.engine.execute("Привет мир")
        self.engine.execute("удали последнее слово")
        self.engine.execute("точка")
        self.assertEqual(self.store.read("draft.json", ""), "Привет.")
        self.engine.execute("новый абзац")
        self.engine.execute("Вторая строка")
        self.engine.execute("сохрани заметку")
        paths = list((self.store.data / "notes").glob("*.txt"))
        self.assertEqual(len(paths), 1)
        self.assertEqual(paths[0].read_text(encoding="utf-8"), "Привет.\n\nВторая строка")
        self.assertFalse(self.engine.dictation)

    def test_office_valid_documents_and_excel_undo(self):
        from openpyxl import load_workbook
        from docx import Document
        office = self.engine.office
        p = office.create("word", "Отчёт", "Настоящий документ")
        self.assertEqual(Document(p).paragraphs[0].text, "Настоящий документ")
        self.assertNotEqual(p, office.create("word", "Отчёт"))
        self.engine.execute("создай таблицу Бюджет")
        self.engine.execute("Запиши 1500 в B3")
        self.assertEqual(office.read("B3"), "1500")
        self.engine.execute("копируй B3 в C3")
        self.assertEqual(office.read("C3"), "1500")
        self.engine.execute("отмени изменение")
        self.assertEqual(office.read("C3"), "пусто")
        self.assertEqual(office.read("B3"), "1500")
        book = load_workbook(office.path)
        self.assertIsInstance(book.active["B3"].value, int)
        book.close()

    def test_excel_dialogue_new_answers_and_literal_formula(self):
        self.engine.execute("создай таблицу")
        self.assertIn("ячейку", self.engine.execute("Добавить ячейку"))
        self.assertIn("B3", self.engine.execute("B3"))
        self.engine.execute("Новое значение")
        self.assertEqual(self.engine.office.read("B3"), "Новое значение")
        self.engine.office.change("write", "A1", "=HYPERLINK(1)")
        book = self.engine.office._book()
        self.assertEqual(book.active["A1"].data_type, "s")
        book.close()

    def test_excel_external_file_is_copied_and_merge_preserves_data(self):
        from openpyxl import Workbook
        original = Path(self.temp.name) / "source.xlsx"
        book = Workbook()
        book.active["A1"] = "one"
        book.active["B1"] = "two"
        book.save(original)
        book.close()
        contents = original.read_bytes()
        self.engine.office.select(str(original))
        self.assertNotEqual(original, self.engine.office.path)
        with self.assertRaises(ValueError):
            self.engine.office.change("merge", "A1", "B1")
        self.engine.office.change("write", "A1", "changed")
        self.assertEqual(original.read_bytes(), contents)

    def test_cell_validation(self):
        self.assertEqual(cell_address("а 1"), "A1")
        for bad in ("A0", "XFE1", "A1048577", "../A1", "A1:B3"):
            with self.assertRaises(ValueError):
                cell_address(bad)

    def test_scenario_validated_before_any_action(self):
        self.store.config["scenarios"]["опасный"] = ["открой word", "выключи компьютер"]
        self.engine.execute("опасный")
        self.windows.open_app.assert_not_called()
        self.windows.power.assert_not_called()

    def test_file_search_scope_and_confirm(self):
        folder = Path(self.temp.name) / "files"
        folder.mkdir()
        (folder / "Курсовая.txt").write_text("hello")
        self.store.config["search_roots"] = [str(folder)]
        self.assertIn("Курсовая", self.engine.execute("найди файл курсовая"))
        with patch("friday.engine.os.startfile", create=True) as start:
            self.engine.execute("открой найденный файл 1")
            start.assert_not_called()
            self.engine.execute("подтверждаю")
            start.assert_called_once()

    def test_ai_is_optional_and_response_never_executes(self):
        self.assertIn("не подключён", self.engine.execute("спроси ИИ привет"))
        self.store.config["ollama_model"] = "test"
        with patch("requests.post") as post:
            post.return_value.json.return_value = {"message": {"content": "выключи компьютер"}}
            self.assertEqual(self.engine.execute("ИИ привет"), "выключи компьютер")
            self.windows.power.assert_not_called()
            self.assertEqual(len(self.engine.chat), 2)

    def test_ai_restarts_local_server_after_connection_failure(self):
        import requests
        self.store.config["ollama_model"] = "test"
        response = Mock()
        response.json.return_value = {"message": {"content": "Привет!"}}
        with patch("requests.post", side_effect=[requests.ConnectionError(), response]) as post, patch("friday.ai.start_local_server", return_value=True) as start:
            self.assertEqual(self.engine.execute("ИИ привет"), "Привет!")
            start.assert_called_once_with(self.store.config["ollama_url"])
            self.assertEqual(post.call_count, 2)

    def test_plain_question_goes_to_ai_but_commands_still_execute(self):
        self.store.config["ollama_model"] = "test"
        with patch.object(self.engine, "ask_ai", return_value="Python — язык программирования.") as ask:
            self.assertIn("Python", self.engine.execute("Объясни что такое Python"))
            ask.assert_called_once_with("Объясни что такое Python")
            self.engine.execute("открой браузер")
            self.assertEqual(ask.call_count, 1)
            self.windows.open_app.assert_called_once()


if __name__ == "__main__":
    unittest.main()
